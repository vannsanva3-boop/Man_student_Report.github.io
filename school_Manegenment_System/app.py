import os
import io
import urllib.parse
from flask import Flask, render_template, request, jsonify, send_file, session, redirect, url_for
from database import get_db, init_db

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as ExcelImage

app = Flask(__name__)
app.secret_key = "grade10_math_tongpo_secret_2026"

def get_grade(total):
    if total >= 95: return "A"
    if total >= 90: return "B"
    if total >= 80: return "C"
    if total >= 65: return "D"
    if total >= 50: return "E"
    return "F"

@app.before_request
def check_auth():
    if request.path.startswith('/static') or request.path in ['/login']:
        return None
    if 'user' not in session:
        if request.path.startswith('/api/'):
            return jsonify({"error": "Unauthorized"}), 401
        return redirect(url_for('login'))

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "GET":
        if 'user' in session: return redirect(url_for('index'))
        return render_template("login.html")

    username = request.form.get("username", "").strip()
    password = request.form.get("password", "").strip()
    db = get_db()
    user = db.execute("SELECT * FROM users WHERE username = ? AND password = ?", (username, password)).fetchone()
    if user:
        session['user'] = user['username']
        session['fullname'] = user['fullname']
        return redirect(url_for('index'))
    return render_template("login.html", error="ឈ្មោះអ្នកប្រើ ឬពាក្យសម្ងាត់មិនត្រឹមត្រូវ!")

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route("/")
def index():
    return render_template("index.html")

# ----------------- APIs ----------------- #

@app.get("/api/settings")
def get_settings():
    db = get_db()
    st = db.execute("SELECT * FROM settings WHERE id = 1").fetchone()
    return jsonify(dict(st) if st else {})

@app.post("/api/settings")
def save_settings():
    try:
        db = get_db()
        d = request.json or {}
        db.execute("""
            UPDATE settings SET
                teacher_name = ?,
                teacher_date = ?,
                admin_lunar_date = ?,
                admin_solar_date = ?,
                academic_year = ?
            WHERE id = 1
        """, (
            d.get("teacher_name", "វ៉ាន់ ជីវ៉ា"),
            d.get("teacher_date", "ថ្ងៃទី ៣០ ខែ កញ្ញា ឆ្នាំ២០២៦"),
            d.get("admin_lunar_date", "ថ្ងៃ  ពុធ ៤រោច ខែ  ភទ្របទ  ឆ្នាំ  មមី  អដ្ឋសក័ ព.ស.២៥៧០"),
            d.get("admin_solar_date", "ភ្នំពេញ ត្រូវនឹងថ្ងៃទី  ៣០ ខែ  កញ្ញា  ឆ្នាំ២០២៦"),
            d.get("academic_year", "មុខវិជ្ជាសិក្សា៖ គណិតវិទ្យា /កម្រិតថ្នាក់ ១០ /ក្នុងឆ្នាំ២០២៦")
        ))
        db.commit()
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.get("/api/dashboard-stats")
def dashboard_stats():
    db = get_db()
    total = db.execute("SELECT COUNT(*) FROM students").fetchone()[0]
    female = db.execute("SELECT COUNT(*) FROM students WHERE gender = 'ស្រី'").fetchone()[0]
    male = db.execute("SELECT COUNT(*) FROM students WHERE gender = 'ប្រុស'").fetchone()[0]
    return jsonify({"total": total, "female": female, "male": male})

@app.get("/api/students")
def get_students():
    db = get_db()
    search = f"%{request.args.get('q', '').strip()}%"

    query = """
        SELECT s.*, 
               COALESCE(sc.behavior, 0) as behavior,
               COALESCE(sc.attendance, 0) as attendance,
               COALESCE(sc.homework, 0) as homework,
               COALESCE(sc.worksheet, 0) as worksheet,
               COALESCE(sc.quiz, 0) as quiz,
               (COALESCE(sc.behavior, 0) + COALESCE(sc.attendance, 0) + COALESCE(sc.homework, 0) + COALESCE(sc.worksheet, 0) + COALESCE(sc.quiz, 0)) as total_score
        FROM students s
        LEFT JOIN math_scores sc ON s.id = sc.student_id
        WHERE (s.name LIKE ? OR s.student_code LIKE ?)
        ORDER BY total_score DESC
    """
    rows = db.execute(query, (search, search)).fetchall()

    students_list = []
    for idx, r in enumerate(rows):
        st = dict(r)
        total = round(st["total_score"], 1)
        avg = round(total / 6.0, 2)

        st["total_score"] = int(total) if total.is_integer() else total
        st["average"] = f"{avg:.2f}"
        st["rank"] = idx + 1
        st["result"] = "ជាប់" if total >= 50 else "ធ្លាក់"
        st["grade"] = get_grade(total)
        students_list.append(st)

    return jsonify({"data": students_list})

@app.post("/api/students/<int:sid>/scores")
def update_scores(sid):
    try:
        db = get_db()
        d = request.json or {}
        db.execute("""
            INSERT INTO math_scores (student_id, behavior, attendance, homework, worksheet, quiz)
            VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(student_id) DO UPDATE SET
                behavior = excluded.behavior,
                attendance = excluded.attendance,
                homework = excluded.homework,
                worksheet = excluded.worksheet,
                quiz = excluded.quiz
        """, (sid, float(d.get("behavior", 0)), float(d.get("attendance", 0)), 
              float(d.get("homework", 0)), float(d.get("worksheet", 0)), float(d.get("quiz", 0))))
        db.commit()
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.post("/api/students")
def add_student():
    db = get_db()
    name = request.form.get("name", "").strip()
    gender = request.form.get("gender", "ស្រី")
    if not name: return jsonify({"error": "សូមបញ្ចូលឈ្មោះ"}), 400

    last_id = db.execute("SELECT MAX(id) FROM students").fetchone()[0] or 0
    student_code = f"STU-{last_id + 1:02d}"

    cursor = db.cursor()
    cursor.execute("INSERT INTO students (student_code, name, gender) VALUES (?, ?, ?)", (student_code, name, gender))
    new_sid = cursor.lastrowid
    cursor.execute("INSERT INTO math_scores (student_id) VALUES (?)", (new_sid,))
    db.commit()
    return jsonify({"success": True}), 201

@app.delete("/api/students/<int:sid>")
def delete_student(sid):
    db = get_db()
    db.execute("DELETE FROM students WHERE id = ?", (sid,))
    db.commit()
    return jsonify({"success": True})

# ----------------- EXPORT EXCEL (សម្រាប់ iPhone & PC មិនឱ្យចេញ .webarchive) ----------------- #
@app.get("/export/excel")
def export_real_excel():
    custom_name = request.args.get("filename", "").strip()
    if not custom_name:
        custom_name = "របាយការណ៍_លទ្ធផលប្រឡងប្រចាំខែ_ថ្នាក់ទី១០.xlsx"
    
    if not custom_name.lower().endswith(".xlsx"):
        custom_name += ".xlsx"

    db = get_db()
    settings = db.execute("SELECT * FROM settings WHERE id = 1").fetchone()
    
    students = db.execute("""
        SELECT s.student_code, s.name, s.gender,
               COALESCE(sc.behavior, 0) as behavior,
               COALESCE(sc.attendance, 0) as attendance,
               COALESCE(sc.homework, 0) as homework,
               COALESCE(sc.worksheet, 0) as worksheet,
               COALESCE(sc.quiz, 0) as quiz,
               (COALESCE(sc.behavior, 0) + COALESCE(sc.attendance, 0) + COALESCE(sc.homework, 0) + COALESCE(sc.worksheet, 0) + COALESCE(sc.quiz, 0)) as total
        FROM students s
        LEFT JOIN math_scores sc ON s.id = sc.student_id
        ORDER BY total DESC
    """).fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "លទ្ធផលប្រចាំខែ"
    
    ws.views.sheetView[0].showGridLines = True
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    col_widths = {
        'A': 3,     'B': 5.5,   'C': 22,    'D': 7.5,
        'E': 9.5,   'F': 9.5,   'G': 11,    'H': 9,
        'I': 11.5,  'J': 10,    'K': 10.5,  'L': 11,
        'M': 9,     'N': 8
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[7].height = 22
    ws.row_dimensions[8].height = 20
    ws.row_dimensions[9].height = 22
    ws.row_dimensions[10].height = 22

    font_khmer = Font(name="Khmer OS Siemreap", size=9.5, bold=False, color="000000")
    font_header_muol = Font(name="Khmer OS Muol Light", size=9, bold=False, color="000000")
    font_title_muol = Font(name="Khmer OS Muol Light", size=9, bold=False, color="000000")
    font_title_muols = Font(name="Khmer OS Muol Light", size=9, bold=False, color="FFFFFF")
    font_red = Font(name="Khmer OS Siemreap", size=9.5, bold=True, color="C00000")
    font_white = Font(name="Khmer OS Siemreap", size=9.5, bold=True, color="FFFFFF")

    fill_orange_header = PatternFill(start_color="C55A11", end_color="C55A11", fill_type="solid")
    fill_cream = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    fill_stat_orange = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
    fill_stat_green = PatternFill(start_color="385723", end_color="385723", fill_type="solid")
    fill_stat_blue = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='7F7F7F'),
        right=Side(style='thin', color='7F7F7F'),
        top=Side(style='thin', color='7F7F7F'),
        bottom=Side(style='thin', color='7F7F7F')
    )

    def apply_border_box(min_col, min_row, max_col, max_row):
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                ws.cell(row=r, column=c).border = thin_border

    logo_path = os.path.join("static", "logo.png")
    if os.path.exists(logo_path):
        try:
            img = ExcelImage(logo_path)
            img.width = 72
            img.height = 72
            ws.add_image(img, 'C2')
        except Exception:
            pass

    ws.merge_cells('E2:K2')
    ws['E2'] = "ព្រះរាជាណាចក្រកម្ពុជា"
    ws['E2'].font = font_header_muol
    ws['E2'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

    ws.merge_cells('E3:K3')
    ws['E3'] = "ជាតិ សាសនា ព្រះមហាក្សត្រ"
    ws['E3'].font = font_header_muol
    ws['E3'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

    divider_path = os.path.join("static", "divider.png")
    if os.path.exists(divider_path):
        try:
            d_img = ExcelImage(divider_path)
            d_img.width = 65
            d_img.height = 16
            ws.add_image(d_img, 'H4')
        except Exception:
            pass

    ws['B5'] = settings['school_name'] if settings and settings['school_name'] else "សាលាបង្វឹក និង បង្រៀនគួរពិសេស តុងប៉"
    ws['B5'].font = font_title_muol

    ws['B6'] = settings['school_branch'] if settings and settings['school_branch'] else "អគារ G   (សាខា ស្វាយប៉ាក)"
    ws['B6'].font = font_title_muol

    ws.merge_cells('E7:K7')
    ws['E7'] = "របាយការណ៍លទ្ធផលប្រឡងប្រចាំខែ"
    ws['E7'].font = font_title_muol
    ws['E7'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

    ws.merge_cells('E8:K8')
    ws['E8'] = settings['academic_year'] if settings and settings['academic_year'] else "មុខវិជ្ជាសិក្សា៖ គណិតវិទ្យា /កម្រិតថ្នាក់ ១០ /ក្នុងឆ្នាំ២០២៦"
    ws['E8'].font = font_title_muol
    ws['E8'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

    ws.merge_cells('B9:D9')
    ws['B9'] = "របាយការណ៍ពិន្ទុគិតជា %"
    for col_c in ['B9', 'C9', 'D9']:
        ws[col_c].fill = fill_orange_header
        ws[col_c].font = font_title_muols
        ws[col_c].alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

    pct_headers = {'E9': "10%", 'F9': "10%", 'G9': "10%", 'H9': "50%", 'I9': "20%", 'J9': "100%"}
    for cid, val in pct_headers.items():
        ws[cid] = val
        ws[cid].fill = fill_orange_header
        ws[cid].font = font_white
        ws[cid].alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

    for col_char, text in [('K', "មធ្យមភាគ"), ('L', "ចំណាត់ថ្នាក់"), ('M', "លទ្ធផល"), ('N', "និទ្ទេស")]:
        ws.merge_cells(f'{col_char}9:{col_char}10')
        ws[f'{col_char}9'] = text
        for r in [9, 10]:
            ws[f'{col_char}{r}'].fill = fill_cream
            ws[f'{col_char}{r}'].font = font_title_muol
            ws[f'{col_char}{r}'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

    for cid, text in [('B10', "ល.រ"), ('C10', "ឈ្មោះសិស្ស"), ('D10', "ភេទ")]:
        ws[cid] = text
        ws[cid].fill = fill_cream
        ws[cid].font = font_title_muol
        ws[cid].alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

    sub_titles = {
        'E10': "ឥរិយាបថ", 'F10': "វត្តមាន", 'G10': "កិច្ចការផ្ទះ",
        'H10': "គណិត", 'I10': "ឆ្លើយសំណួរ", 'J10': "ពិន្ទុសរុប"
    }
    for cid, val in sub_titles.items():
        ws[cid] = val
        ws[cid].fill = fill_cream
        ws[cid].font = font_title_muol
        ws[cid].alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

    apply_border_box(min_col=2, min_row=9, max_col=14, max_row=10)

    current_row = 11
    for idx, s in enumerate(students):
        tot = round(s["total"], 1) if s["total"] is not None else 0
        avg = round(tot / 6.0, 2)
        res = "ជាប់" if tot >= 50 else "ធ្លាក់"
        g = get_grade(tot)

        ws.row_dimensions[current_row].height = 19

        cell_no = ws.cell(row=current_row, column=2, value=idx + 1)
        cell_no.font = font_khmer
        cell_no.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

        cell_name = ws.cell(row=current_row, column=3, value=s["name"])
        cell_name.font = font_khmer
        cell_name.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

        cell_gen = ws.cell(row=current_row, column=4, value=s["gender"])
        cell_gen.font = font_khmer
        cell_gen.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

        scores = [s["behavior"], s["attendance"], s["homework"], s["worksheet"], s["quiz"]]
        for col_i, sc_val in enumerate(scores, start=5):
            c = ws.cell(row=current_row, column=col_i, value=sc_val)
            c.font = font_khmer
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

        c_tot = ws.cell(row=current_row, column=10, value=int(tot) if float(tot).is_integer() else tot)
        c_tot.font = font_red
        c_tot.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

        c_avg = ws.cell(row=current_row, column=11, value=f"{avg:.2f}")
        c_avg.font = font_red
        c_avg.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

        c_rank = ws.cell(row=current_row, column=12, value=idx + 1)
        c_rank.font = font_red
        c_rank.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

        c_res = ws.cell(row=current_row, column=13, value=res)
        c_res.font = font_red
        c_res.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

        c_grd = ws.cell(row=current_row, column=14, value=g)
        c_grd.font = font_red
        c_grd.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

        for col_idx in range(2, 15):
            ws.cell(row=current_row, column=col_idx).border = thin_border

        current_row += 1

    stat_start_row = current_row + 2

    ws.merge_cells(start_row=stat_start_row, start_column=2, end_row=stat_start_row+2, end_column=2)
    stat_lbl = ws.cell(row=stat_start_row, column=2, value="គិតជាភេទ")
    stat_lbl.font = font_white
    stat_lbl.alignment = Alignment(horizontal='center', vertical='center', text_rotation=90, wrap_text=False)
    for r in range(stat_start_row, stat_start_row+3):
        ws.cell(row=r, column=2).fill = fill_orange_header

    female_cnt = sum(1 for s in students if s["gender"] == "ស្រី")
    male_cnt = sum(1 for s in students if s["gender"] == "ប្រុស")

    stats_data = [
        ("សិស្សសរុបចំនួន", len(students), fill_stat_orange),
        ("ស្រី", female_cnt, fill_stat_green),
        ("ប្រុស", male_cnt, fill_stat_blue)
    ]

    for offset, (lbl, cnt, fill_c) in enumerate(stats_data):
        r_idx = stat_start_row + offset
        c_lbl = ws.cell(row=r_idx, column=3, value=lbl)
        c_lbl.fill = fill_c
        c_lbl.font = font_white
        c_lbl.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

        c_val = ws.cell(row=r_idx, column=4, value=cnt)
        c_val.fill = fill_c
        c_val.font = font_white
        c_val.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

        c_unit = ws.cell(row=r_idx, column=5, value="នាក់")
        c_unit.fill = fill_c
        c_unit.font = font_white
        c_unit.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)

    apply_border_box(min_col=2, min_row=stat_start_row, max_col=5, max_row=stat_start_row+2)

    lunar_date = settings['admin_lunar_date'] if settings and settings['admin_lunar_date'] else "ថ្ងៃ  ពុធ ៤រោច ខែ  ភទ្របទ  ឆ្នាំ  មមី  អដ្ឋសក័ ព.ស.២៥៧០"
    solar_date = settings['admin_solar_date'] if settings and settings['admin_solar_date'] else "ភ្នំពេញ ត្រូវនឹងថ្ងៃទី  ៣០ ខែ  កញ្ញា  ឆ្នាំ២០២៦"
    t_date = settings['teacher_date'] if settings and settings['teacher_date'] else "ថ្ងៃទី ៣០ ខែ កញ្ញា ឆ្នាំ២០២៦"
    t_name = settings['teacher_name'] if settings and settings['teacher_name'] else "វ៉ាន់ ជីវ៉ា"

    ws.merge_cells(start_row=stat_start_row, start_column=8, end_row=stat_start_row, end_column=13)
    ws.cell(row=stat_start_row, column=8, value=lunar_date).alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
    ws.cell(row=stat_start_row, column=8).font = font_khmer

    ws.merge_cells(start_row=stat_start_row+1, start_column=8, end_row=stat_start_row+1, end_column=13)
    ws.cell(row=stat_start_row+1, column=8, value=solar_date).alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
    ws.cell(row=stat_start_row+1, column=8).font = font_khmer

    ws.merge_cells(start_row=stat_start_row+2, start_column=9, end_row=stat_start_row+2, end_column=12)
    ws.cell(row=stat_start_row+2, column=9, value="បានឃើញ និងឯកភាព").alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
    ws.cell(row=stat_start_row+2, column=9).font = font_title_muol

    ws.merge_cells(start_row=stat_start_row+3, start_column=9, end_row=stat_start_row+3, end_column=12)
    ws.cell(row=stat_start_row+3, column=9, value="គណ:គ្រប់គ្រងសាខា ស្វាយប៉ាក").alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
    ws.cell(row=stat_start_row+3, column=9).font = font_title_muol

    ws.merge_cells(start_row=stat_start_row+3, start_column=2, end_row=stat_start_row+3, end_column=5)
    ws.cell(row=stat_start_row+3, column=2, value=t_date).alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
    ws.cell(row=stat_start_row+3, column=2).font = font_khmer

    ws.merge_cells(start_row=stat_start_row+4, start_column=2, end_row=stat_start_row+4, end_column=5)
    ws.cell(row=stat_start_row+4, column=2, value="គ្រូទទួលបន្ទុកថ្នាក់").alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
    ws.cell(row=stat_start_row+4, column=2).font = font_title_muol

    ws.merge_cells(start_row=stat_start_row+6, start_column=2, end_row=stat_start_row+6, end_column=5)
    ws.cell(row=stat_start_row+6, column=2, value=t_name).alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
    ws.cell(row=stat_start_row+6, column=2).font = font_title_muol

    output = io.BytesIO()
    wb.save(output)
    excel_bytes = output.getvalue()
    output.seek(0)

    encoded_name = urllib.parse.quote(custom_name.encode('utf-8'))

    # បង្ខំ Header ឱ្យ iPhone ស្គាល់ជាឯកសារ Excel ពិតប្រាកដ
    response = send_file(
        io.BytesIO(excel_bytes),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=custom_name
    )

    response.headers["Content-Disposition"] = f"attachment; filename=\"report.xlsx\"; filename*=UTF-8''{encoded_name}"
    response.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    response.headers["Content-Length"] = str(len(excel_bytes))
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"

    return response

if __name__ == "__main__":
    init_db()
    app.run(debug=True, host="0.0.0.0", port=5000)